import openpyxl
from openpyxl import Workbook
from typing import Tuple, List
from langchain.prompts import PromptTemplate
from langchain_openai import ChatOpenAI
from dotenv import load_dotenv


load_dotenv(override=True)


class ScreeningDecision:
    def __init__(self, model_name: str = "o1-mini", temperature: float = 0.0):
        self.llm = ChatOpenAI(model_name=model_name, temperature=temperature)

        # Define a prompt template that requires two variables: question, answer
        self.yes_no_prompt_template = PromptTemplate(
            input_variables=["question", "answer"],
            template="""
            The user question was: {question}
            The LLM's answer was: {answer}

            Based on that answer, respond with a single word:
            YES
            NO
            """,
        )

    def decide_inclusion(self, question: str, answer: str) -> Tuple[str, str]:
        """
        Use the new chain style to decide if we say 'YES' or 'NO'.
        Explanation is the entire 'answer' text.
        """
        # Create a small chain by piping the PromptTemplate into the LLM
        sequence = self.yes_no_prompt_template | self.llm

        # Invoke with the two required variables
        response = sequence.invoke({"question": question, "answer": answer})

        # response.content will have the model's raw text
        # We strip and uppercase to interpret the single word
        decision_raw = response.content.strip().upper()
        if "YES" in decision_raw:
            decision = "YES"
        else:
            decision = "NO"

        # Explanation is the entire Q/A answer
        explanation = answer.strip()
        return (decision, explanation)

    def write_decisions_for_article(
        self, excel_path: str, article_name: str, decisions: List[Tuple[str, str]]
    ):
        """
        Create a new sheet for the article_name,
        and write each (YES/NO, explanation) pair on a new row.
        No header row.
        """
        try:
            wb = openpyxl.load_workbook(excel_path)
        except FileNotFoundError:
            wb = Workbook()

        # Create a new sheet named after the article
        sheet_name = article_name[:31]  # Excel limit
        ws = wb.create_sheet(title=sheet_name)

        row_idx = 1
        # Write row by row (YES/NO, explanation)
        for decision, explanation in decisions:
            ws.cell(row=row_idx, column=1, value=decision)

            # color line red if it is "NO", green if it is "YES"
            fill_color = openpyxl.styles.PatternFill(
                start_color="FF0000" if decision == "NO" else "00FF00",
                end_color="FF0000" if decision == "NO" else "00FF00",
                fill_type="solid"
            )
            for col_idx in range(1, ws.max_column + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill_color
            ws.cell(row=row_idx, column=2, value=explanation)
            row_idx += 1

        # Remove default "Sheet" if it's still empty
        if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
            default_sheet = wb["Sheet"]
            if (
                default_sheet.max_row == 1
                and default_sheet.max_column == 1
                and not default_sheet.cell(1, 1).value
            ):
                wb.remove(default_sheet)

        wb.save(excel_path)
